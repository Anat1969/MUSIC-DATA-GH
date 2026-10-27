# -*- coding: utf-8 -*-
"""קבצי Excel לדוגמה, להורדה מהאפליקציה ולגרירה חזרה אליה.
כל המספרים הם פרופילים מייצגים (משוערים): נכונים בצורתם ובהתנהגותם, לא נשלפו
ממקור רשמי. זה כתוב בכל קובץ, בגיליון "הסבר" ובשורת המקור.
העמודה המספרית הראשונה היא המנגינה, השנייה ההרמוניה. לכן הסדר שלהן מכוון.

הרצה מתיקיית השורש: python3 tools/build_samples.py
"""
import os, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'samples')
NOTE = 'פרופיל מייצג (משוער). לא נתון רשמי. להחלפה בנתוני העירייה או הלמ"ס לפני שימוש תכנוני.'

def book(path, title, headers, rows, story, listen, settings):
    wb = Workbook(); ws = wb.active; ws.title = 'נתונים'
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for r in rows: ws.append(r)
    head = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='16181B')
    thin = Side(style='thin', color='D0D0D0')
    for c in ws[1]:
        c.font = head; c.fill = fill; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = Border(bottom=thin)
            c.alignment = Alignment(horizontal='center')
    ws.row_dimensions[1].height = 34
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(str(h)) + 4)
    ws.freeze_panes = 'B2'

    ex = wb.create_sheet('הסבר'); ex.sheet_view.rightToLeft = True
    lines = [('כותרת', title), ('מה הנתונים מספרים', story), ('מה תשמעו', listen),
             ('הגדרה מומלצת', settings), ('מקור', NOTE),
             ('איך משתמשים', 'גוררים את הקובץ אל "הנתונים שלך" ב-DataSong. העמודה הראשונה היא התוויות, '
                             'העמודה המספרית הראשונה היא המנגינה והשנייה ההרמוניה. אפשר לשנות כל עמודה ולגרור שוב.')]
    for k, v in lines: ex.append([k, v])
    ex.column_dimensions['A'].width = 20; ex.column_dimensions['B'].width = 110
    for row in ex.iter_rows():
        row[0].font = Font(bold=True); row[1].alignment = Alignment(wrap_text=True, vertical='top')
    os.makedirs(OUT, exist_ok=True)
    wb.save(os.path.join(OUT, path))
    print('built', path, len(rows), 'rows')

# ---------------------------------------------------------------- 1. מבני חינוך לפי רובע
# הרובעים לפי סדר האכלוס, מהוותיק לחדש. הסדר הוא ציר הזמן של העיר.
quarters = ['א','ב','ג','ד','ה','ו','ז','ח','ט','י','י"א','י"ב','ט"ו','ט"ז','י"ז']
years    = [1958,1961,1964,1967,1971,1975,1980,1985,1990,1993,1996,2002,2006,2012,2018]
pop      = [12.5,14.0,15.5,11.0,16.0,13.5,22.0,17.5,16.5,13.0,12.0,15.0,24.0,18.0,14.0]       # אלפי תושבים
# גנים: רובעים חדשים ומשפחות צעירות, ושיא ברובעים עם משפחות גדולות (ז', ט"ו)
kg       = [9,10,12,8,13,12,34,18,15,12,12,17,38,30,26]
prim     = [3,3,4,3,4,4,9,5,5,4,4,5,10,7,5]
sec      = [3,2,3,2,3,2,4,3,3,2,2,3,4,2,1]
rows=[]
for q,y,p,k,pr,s in zip(quarters,years,pop,kg,prim,sec):
    tot=k+pr+s
    rows.append([f'רובע {q}', round(tot/p,2), k, pr, s, tot, p, y])
book('01_ashdod_education_by_quarter.xlsx', 'מבני חינוך לפי רובע באשדוד',
     ['רובע','מבני חינוך ל-1,000 תושבים','גני ילדים','בתי ספר יסודיים','על-יסודי','סה"כ מבני חינוך','תושבים (אלפים)','שנת אכלוס'],
     rows,
     'הרובעים מסודרים לפי שנת האכלוס. הוותיקים נשענים על על-יסודי ומעט גנים; רובעים עם משפחות גדולות קופצים בגנים; '
     'החדשים ביותר נבנו סביב גיל הגן. הצפיפות של מבני החינוך היא מראה של מבנה המשפחה, לא רק של גודל הרובע.',
     'המנגינה היא מבני חינוך לאלף תושבים: שקטה בוותיקים, קופצת ברובעים של משפחות גדולות, ועולה שוב בחדשים. '
     'ההרמוניה נשענת על הגנים, ולכן היא משנה אופי בדיוק במקום שבו הרובע "צעיר".',
     'סולם דוריאני, מנגינה מרימבה, הרמוניה כרית מיתרים, טיפולוגיה: חזית.')

# ---------------------------------------------------------------- 2. שימושי קרקע לאורך חתך
# חתך ממערב למזרח, 24 קטעים של כ-400 מ'. בכל קטע: אחוז השטח לכל שימוש (סך הכול 100).
zones = (['חוף ומרינה']*3 + ['רובע י"א']*3 + ['מרכז העיר (סיטי)']*4 + ['רובע ט\'']*3 +
         ['רובע ח\'']*3 + ['רובע ג\'']*3 + ['אזור תעסוקה צפוני']*3 + ['שולי העיר']*2)
def mix(i):
    t=i/23
    leisure = 55*math.exp(-((t-0.02)/0.10)**2) + 14*math.exp(-((t-0.62)/0.05)**2) + 6
    commerce= 8 + 34*math.exp(-((t-0.36)/0.08)**2) + 6*math.exp(-((t-0.05)/0.05)**2)
    employ  = 6 + 52*math.exp(-((t-0.86)/0.08)**2) + 12*math.exp(-((t-0.38)/0.07)**2)
    resid   = max(5, 100 - leisure - commerce - employ)
    s = resid+employ+commerce+leisure
    return [round(100*v/s,1) for v in (resid,employ,commerce,leisure)]
rows=[]
for i,z in enumerate(zones):
    r,e,c,l = mix(i)
    rows.append([z if (i==0 or zones[i-1]!=z) else '', r, e, c, l, round(i*0.4,1)])
book('02_ashdod_land_use_mix.xlsx', 'תמהיל שימושי קרקע לאורך חתך עירוני',
     ['אזור','מגורים %','תעסוקה %','מסחר %','נופש ופנאי %','מרחק מהים (ק"מ)'],
     rows,
     'חתך מהים מזרחה. בחוף שולט הפנאי, במרכז העיר המסחר והתעסוקה נכנסים לתוך המגורים, ברובעים המגורים מובילים, '
     'ובמזרח התעסוקה תופסת את רוב השטח. היחס בין ארבעת השימושים הוא מה שמגדיר אם רחוב חי כל היום או רק בערב.',
     'המנגינה היא המגורים: שקע בחוף, רמה ארוכה ברובעים, ונפילה באזור התעסוקה. '
     'ההרמוניה היא התעסוקה, שעולה כשהמגורים יורדים. שתי שכבות שזזות הפוך נשמעות כמו דיאלוג. '
     'כדאי להעביר את הביס למסחר ואת הקצב לנופש ולשמוע איך מרכז העיר מתמלא.',
     'סולם מז\'ורי, מנגינה פסנתר, הרמוניה מקהלה, ביס קונטרבס, טיפולוגיה: חתך.')

# ---------------------------------------------------------------- 3. שכר לפי גיל וענף
ages = list(range(22, 68, 2))
def curve(a, start, peak, peak_age, late, width):
    # עלייה עם הוותק עד שיא, ואחריו התמתנות או ירידה קלה
    if a <= peak_age:
        x=(a-22)/max(1,(peak_age-22)); return start + (peak-start)*(1-(1-x)**1.8)
    return peak + (late-peak)*min(1,((a-peak_age)/width))
sectors = {
    'הייטק':           (14, 46, 44, 38, 16),
    'שירות ציבורי':     (8,  22, 60, 22, 4),
    'תעשייה':          (8.5, 18, 48, 16, 14),
    'מסחר וקמעונאות':   (6.5, 11, 42, 10, 20),
    'חינוך':           (7.5, 19, 58, 19, 6),
    'בריאות':          (9,  26, 54, 24, 10),
}
rows=[]
for a in ages:
    rows.append([f'גיל {a}'] + [round(curve(a,*sectors[k]),1) for k in sectors])
book('03_wage_by_age_and_sector.xlsx', 'שכר חודשי ממוצע לפי גיל וענף',
     ['גיל'] + [f'{k} (אלפי ₪)' for k in sectors],
     rows,
     'שש עקומות שכר לאורך חיי העבודה. בהייטק השכר מזנק מוקדם ומגיע לשיא באמצע הארבעים; בשירות הציבורי ובחינוך הוא '
     'עולה לאט ובהתמדה עם הוותק; במסחר העקומה כמעט שטוחה. הפער בין הצורות חשוב לא פחות מהפער בגובה.',
     'המנגינה היא הייטק: טיפוס מהיר ושיא מוקדם. ההרמוניה היא השירות הציבורי, שעדיין עולה כשההייטק כבר התמתן. '
     'נסו לשים מסחר במנגינה: קו כמעט ישר, והקצב כמעט שותק, כי אין נקודות מפנה.',
     'סולם מינורי, מנגינה צ\'לו, הרמוניה כרית מיתרים, ביס דרון, טיפולוגיה: מסה.')
